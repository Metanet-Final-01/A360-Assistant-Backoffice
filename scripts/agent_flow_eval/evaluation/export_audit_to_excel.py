"""audit_final_goldset.py가 만든 audit.json을 사람이 검수하기 쉬운 다중 시트
엑셀로 정리한다. 요약 숫자만이 아니라 매칭 쌍/미매칭/Judge 판정 로그/핵심업무
분류 로그/분기 진단까지 세부사항 전부를 시트로 나눠 담는다."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 팔레트 - 차분한 남색/청회색 계열(강조색 하나만 씀, 임의 화려한 색 안 씀)
_INK = "1F2937"
_MUTED = "6B7280"
_RULE_FILL = "1E3A5F"
_JUDGE_FILL = "2E5C7E"
_WORF_FILL = "8B95A1"
_HEADER_TEXT = "FFFFFF"
_ROW_ALT_FILL = "F3F4F6"
_BORDER_COLOR = "D1D5DB"
_ACCENT = "1E3A5F"


def load_audit(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def build_overview_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    overview = []
    for row in rows:
        score = row["action_prf1"]
        rule_score = row["rule_only_action_prf1"]
        chain = row["action_chain"]
        rule_chain = row["rule_only_action_chain"]
        branch = row.get("branch_coverage") or {}
        core_cls = row.get("core_business_classification") or {}
        worfbench = row.get("worfbench") or {}
        overview.append(
            {
                "case_id": row["case_id"],
                "version": row["version"],
                "sample": row.get("sample"),
                "run_id": row["run_id"],
                "gold_count_core": score["gold_count"],
                "pred_count_core": score["pred_count"],
                "gold_excluded_non_core": core_cls.get("gold_excluded_count"),
                "pred_excluded_non_core": core_cls.get("pred_excluded_count"),
                "rule_only_tp": rule_score["tp"],
                "rule_only_precision": rule_score["precision"],
                "rule_only_recall": rule_score["recall"],
                "rule_only_f1": rule_score["f1"],
                "rule_only_chain_precision": rule_chain["precision"],
                "rule_only_chain_recall": rule_chain["recall"],
                "rule_only_chain_f1_lis": rule_chain["f1"],
                "judge_assisted_tp": score["tp"],
                "judge_assisted_precision": score["precision"],
                "judge_assisted_recall": score["recall"],
                "judge_assisted_f1": score["f1"],
                "judge_assisted_chain_precision": chain["precision"],
                "judge_assisted_chain_recall": chain["recall"],
                "judge_assisted_chain_f1_lis": chain["f1"],
                "worfbench_status_external_reference": worfbench.get("status"),
                "worfbench_precision_external_reference": worfbench.get("precision"),
                "worfbench_recall_external_reference": worfbench.get("recall"),
                "worfbench_f1_external_reference": worfbench.get("f1_score"),
                "branch_applicable": branch.get("applicable"),
                "branch_coverage": branch.get("branch_coverage"),
                "branch_score": branch.get("branch_score"),
                "branch_group_count": branch.get("group_count"),
                "branch_scoreable_count": branch.get("scoreable_branch_count"),
                "control_if_gold": row["control_counts"]["gold"].get("if"),
                "control_loop_gold": row["control_counts"]["gold"].get("loop"),
                "control_try_gold": row["control_counts"]["gold"].get("try"),
                "unmatched_gold_count": len(row["unmatched_gold"]),
                "unmatched_prediction_count": len(row["unmatched_prediction"]),
            }
        )
    return overview


def build_matches_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        for m in row["matches"]:
            out.append(
                {
                    "case_id": row["case_id"],
                    "version": row["version"],
                    "sample": row.get("sample"),
                    "match_type": m["match_type"],
                    "canonical_label": m["canonical_label"],
                    "gold_label": m["gold_label"],
                    "prediction_label": m["prediction_label"],
                    "gold_id": m["gold_id"],
                    "pred_id": m["pred_id"],
                }
            )
    return out


def build_unmatched_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        for item in row["unmatched_gold"]:
            out.append(
                {
                    "case_id": row["case_id"],
                    "version": row["version"],
                    "sample": row.get("sample"),
                    "side": "gold_only(FN)",
                    "label": item["label"],
                    "canonical_label": item["canonical_label"],
                    "uid": item["uid"],
                }
            )
        for item in row["unmatched_prediction"]:
            out.append(
                {
                    "case_id": row["case_id"],
                    "version": row["version"],
                    "sample": row.get("sample"),
                    "side": "prediction_only(FP)",
                    "label": item["label"],
                    "canonical_label": item["canonical_label"],
                    "uid": item["uid"],
                }
            )
    return out


def build_judge_log_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        gold_by_uid = row.get("gold_actions_by_uid") or {}
        pred_by_uid = row.get("pred_actions_by_uid") or {}
        for entry in row["judge_log"]:
            out.append(
                {
                    "case_id": row["case_id"],
                    "version": row["version"],
                    "sample": row.get("sample"),
                    "gold_id": entry.get("gold_id"),
                    "gold_label": (gold_by_uid.get(entry.get("gold_id"), {}) or {}).get("canonical_label"),
                    "pred_id": entry.get("pred_id"),
                    "prediction_label": (pred_by_uid.get(entry.get("pred_id"), {}) or {}).get("canonical_label"),
                    "similarity": entry.get("similarity"),
                    "verdict": entry.get("verdict", entry.get("status")),
                    "reason": entry.get("reason") or entry.get("error"),
                    "stage": entry.get("stage"),
                }
            )
    return out


def build_core_business_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        core_cls = row.get("core_business_classification") or {}
        for side, log_key in (("gold", "gold_log"), ("pred", "pred_log")):
            for entry in core_cls.get(log_key) or []:
                out.append(
                    {
                        "case_id": row["case_id"],
                        "version": row["version"],
                        "sample": row.get("sample"),
                        "side": side,
                        "uid": entry.get("uid"),
                        "label": entry.get("label"),
                        "params": entry.get("params"),
                        "verdict": entry.get("verdict"),
                        "source": entry.get("source"),
                        "reason": entry.get("reason"),
                    }
                )
    return out


def build_branch_detail_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        branch = row.get("branch_coverage") or {}
        for group in branch.get("groups") or []:
            for b in group["branches"]:
                out.append(
                    {
                        "case_id": row["case_id"],
                        "version": row["version"],
                        "sample": row.get("sample"),
                        "group_id": group["group_id"],
                        "branch_name": b["branch_name"],
                        "core_action_count": b["core_action_count"],
                        "matched_count": b["matched_count"],
                        "coverage_ratio": b["coverage_ratio"],
                    }
                )
    return out


def build_aggregates_rows(aggregates: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    """aggregates의 키는 단일 샘플이면 "v2"/"v3"이고, 여러 샘플을 결합했으면
    "v2__repro_20260803" 같은 "버전__샘플" 형식이다 - 둘 다 지원하도록 분리한다."""
    rows = []
    for key, values in aggregates.items():
        if "__" in key:
            version, sample = key.split("__", 1)
        else:
            version, sample = key, None
        rows.append({"version": version, "sample": sample, **values})
    return rows


def _average_by_version(aggregates: dict[str, dict[str, Any]]) -> dict[str, dict[str, float | int | None]]:
    """샘플이 여러 개면(재현성 반복 실행) 버전별로 평균 낸다. 단일 샘플이면
    그 값 그대로. 숫자 지표만 평균 대상 - case_count 등은 합산이 더 맞지만
    이 헤드라인 페이지는 P/R/F1 평균만 보여주므로 case_count는 그대로 둔다."""
    from collections import defaultdict

    by_version: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for key, values in aggregates.items():
        version = key.split("__", 1)[0] if "__" in key else key
        by_version[version].append(values)

    numeric_keys = [
        "macro_rule_only_action_precision", "macro_rule_only_action_recall", "macro_rule_only_action_f1",
        "macro_action_precision", "macro_action_recall", "macro_action_f1",
        "macro_worfbench_precision_external_reference", "macro_worfbench_recall_external_reference",
        "macro_worfbench_f1_external_reference",
    ]
    result: dict[str, dict[str, float | int | None]] = {}
    for version, samples in by_version.items():
        averaged: dict[str, float | int | None] = {"sample_count": len(samples)}
        for key in numeric_keys:
            values = [s[key] for s in samples if s.get(key) is not None]
            averaged[key] = (sum(values) / len(values)) if values else None
        result[version] = averaged
    return result


def write_headline_sheet(workbook, aggregates: dict[str, dict[str, Any]], *, title: str, subtitle: str) -> None:
    """맨 앞 페이지 - Rule-only/Judge-assisted/WorFBench 핵심 P/R/F1만 재현성
    평균으로 깔끔하게 보여준다(PPT에 그대로 옮겨도 될 정도로). 세부 시트는
    뒤쪽 Overview/Aggregates 등에 그대로 남아있다 - 이 페이지는 요약 전용."""
    ws = workbook.create_sheet(title, 0)
    ws.sheet_view.showGridLines = False

    versions_present = ["v1", "v2", "v3"]
    averaged = _average_by_version(aggregates)
    ordered_versions = [v for v in versions_present if v in averaged]

    thin = Side(style="thin", color=_BORDER_COLOR)
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    n_cols = 1 + 3 * 3  # Agent + 3 groups x (P,R,F1)
    last_col_letter = get_column_letter(n_cols)

    # 제목/부제
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Segoe UI", size=20, bold=True, color=_INK)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Segoe UI", size=11, italic=True, color=_MUTED)
    ws.row_dimensions[2].height = 20

    header_row = 4
    group_row = header_row - 1

    ws.row_dimensions[group_row].height = 22
    ws.row_dimensions[header_row].height = 20

    groups = [
        ("Rule-only  (LLM 호출 없음 · 재현 가능한 기준선)", _RULE_FILL, "macro_rule_only_action_precision", "macro_rule_only_action_recall", "macro_rule_only_action_f1"),
        ("Judge-assisted  (Rule + LLM 판정 포함)", _JUDGE_FILL, "macro_action_precision", "macro_action_recall", "macro_action_f1"),
        ("WorFBench  (외부 벤치마크 참고)", _WORF_FILL, "macro_worfbench_precision_external_reference", "macro_worfbench_recall_external_reference", "macro_worfbench_f1_external_reference"),
    ]

    # Agent 라벨 칸(세로 병합)
    ws.merge_cells(f"A{group_row}:A{header_row}")
    agent_cell = ws[f"A{group_row}"]
    agent_cell.value = "Agent"
    agent_cell.font = Font(name="Segoe UI", size=11, bold=True, color=_HEADER_TEXT)
    agent_cell.fill = PatternFill("solid", fgColor=_INK)
    agent_cell.alignment = center
    agent_cell.border = cell_border

    col = 2
    for group_name, fill_color, *_keys in groups:
        start_letter = get_column_letter(col)
        end_letter = get_column_letter(col + 2)
        ws.merge_cells(f"{start_letter}{group_row}:{end_letter}{group_row}")
        gcell = ws[f"{start_letter}{group_row}"]
        gcell.value = group_name
        gcell.font = Font(name="Segoe UI", size=11, bold=True, color=_HEADER_TEXT)
        gcell.fill = PatternFill("solid", fgColor=fill_color)
        gcell.alignment = center
        for c in range(col, col + 3):
            ws.cell(row=group_row, column=c).border = cell_border
        for sub_label, c in zip(("Precision", "Recall", "F1"), range(col, col + 3)):
            hcell = ws.cell(row=header_row, column=c)
            hcell.value = sub_label
            hcell.font = Font(name="Segoe UI", size=10, bold=True, color=_HEADER_TEXT)
            hcell.fill = PatternFill("solid", fgColor=fill_color)
            hcell.alignment = center
            hcell.border = cell_border
        col += 3

    ws.column_dimensions["A"].width = 12
    for c in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    data_row = header_row + 1
    for i, version in enumerate(ordered_versions):
        values = averaged[version]
        row_fill = PatternFill("solid", fgColor=_ROW_ALT_FILL) if i % 2 == 1 else PatternFill(fill_type=None)
        ws.row_dimensions[data_row].height = 22

        label_cell = ws.cell(row=data_row, column=1)
        label_cell.value = version.upper()
        label_cell.font = Font(name="Segoe UI", size=13, bold=True, color=_ACCENT)
        label_cell.alignment = center
        label_cell.fill = row_fill
        label_cell.border = cell_border

        col = 2
        for _group_name, _fill_color, p_key, r_key, f1_key in groups:
            for key in (p_key, r_key, f1_key):
                value = values.get(key)
                cell = ws.cell(row=data_row, column=col)
                cell.value = value if value is not None else "-"
                if isinstance(value, (int, float)):
                    cell.number_format = "0.000"
                is_f1_col = key == f1_key
                cell.font = Font(name="Segoe UI", size=12 if is_f1_col else 11, bold=is_f1_col, color=_INK)
                cell.alignment = center
                cell.fill = row_fill
                cell.border = cell_border
                col += 1
        data_row += 1

    sample_counts = ", ".join(f"{v.upper()} {averaged[v]['sample_count']}회" for v in ordered_versions)
    note_row = data_row + 1
    ws.merge_cells(f"A{note_row}:{last_col_letter}{note_row}")
    note_cell = ws[f"A{note_row}"]
    note_cell.value = (
        f"※ 재현성: 같은 조건으로 반복 실행한 결과의 평균값입니다 ({sample_counts}). "
        "Rule-only는 LLM 호출 없이 재현되는 공식 기준선이고, Judge-assisted/WorFBench는 참고용입니다."
    )
    note_cell.font = Font(name="Segoe UI", size=9, italic=True, color=_MUTED)
    note_cell.alignment = left
    ws.row_dimensions[note_row].height = 26

    ws.freeze_panes = f"A{header_row + 1}"


def build_glossary_rows() -> list[dict[str, Any]]:
    return [
        {"영역": "핵심업무 분류", "지표/컬럼": "gold_count_core / pred_count_core", "의미": "핵심업무로 분류된 액션 개수(범용 셋업 제외 후). 채점 분모/분자에 실제로 쓰이는 값.", "비고": "classify_core_business_actions() - 규칙(키워드) 우선, 남은 것만 gpt-4o-mini 판정. Gold/예측 양쪽에 대칭 적용."},
        {"영역": "핵심업무 분류", "지표/컬럼": "gold_excluded_non_core / pred_excluded_non_core", "의미": "범용 셋업(로그 폴더, 경로조립, 검증 메시지 등)으로 분류돼 채점에서 제외된 개수.", "비고": "Gold만 적용되는 게 아니라 예측에서도 실제로 발생함(비대칭 아님)."},
        {"영역": "Rule-only", "지표/컬럼": "rule_only_precision/recall/f1", "의미": "canonical label 완전일치 + 조건부 동치 규칙만으로 매칭한 Action Precision/Recall/F1. LLM 호출 없음.", "비고": "재현 가능한 공식 기준선. 실행마다 값이 절대 안 바뀜."},
        {"영역": "Rule-only", "지표/컬럼": "rule_only_chain_precision/recall/f1_lis", "의미": "Rule Match 쌍만으로 계산한 Action Chain 지표. 매칭된 쌍을 예측 순서대로 정렬해 정답 인덱스의 최장증가부분수열(LIS) 길이를 구함.", "비고": "\"순서 정확도\"가 아니라 \"액션 선택+상대순서\"를 함께 반영 - 누락/불필요 액션도 분모에 영향."},
        {"영역": "Judge-assisted", "지표/컬럼": "judge_assisted_precision/recall/f1", "의미": "Rule Match + Judge Match(임베딩 상호Top-1 후보만 gpt-4o-mini에 물어봄) 포함 Action P/R/F1.", "비고": "Judge 판정은 완전히 결정적이진 않아 실행마다 소폭 변동 가능 - 그래서 Rule-only를 공식 기준선으로 따로 둠."},
        {"영역": "Judge-assisted", "지표/컬럼": "judge_assisted_chain_precision/recall/f1_lis", "의미": "Rule+Judge 매칭 쌍으로 계산한 Action Chain(LIS) 지표.", "비고": "LIS 결과는 같은 매칭 쌍 위에서 LCS-DP로도 계산해 교차검증함(action_chain.py) - 다르면 매칭 로직 버그."},
        {"영역": "WorFBench(외부 참고)", "지표/컬럼": "worfbench_precision/recall/f1_external_reference", "의미": "원본 WorFEval 벤더 라이브러리(t_eval_nodes)를 그대로 재현한 외부 벤치마크 점수. all-mpnet-base-v2 임베딩, 유사도 임계값 0.6, 실제 networkx 최대가중매칭.", "비고": "우리 action_equivalence_rules나 핵심업무 분류를 적용하지 않은 완전히 별도 경로 - 우리 에이전트를 판단하는 용도가 아니라 공개 벤치마크와 비교하기 위한 참고값."},
        {"영역": "분기 진단(별도)", "지표/컬럼": "branch_coverage / branch_score", "의미": "if/elseIf/else 상호배타적 분기가 flatten 시 한 리스트로 풀려서 생기는 중복카운트 문제에 대한 진단. Branch Coverage=완전매칭 분기 수/전체, Branch Score=분기별 매칭비율 평균.", "비고": "메인 Action P/R/F1/Chain에는 전혀 반영 안 됨. 새 분기매칭 알고리즘 없이 이미 계산된 Rule/Judge 매칭 결과만 재사용."},
        {"영역": "분기 진단(별도)", "지표/컬럼": "branch_applicable", "의미": "이 케이스에 핵심업무 액션을 포함한 if 분기가 있어서 위 두 지표가 의미 있는지 여부.", "비고": "false면 branch_coverage/branch_score는 공란(억지로 0점 처리 안 함)."},
        {"영역": "재현성", "지표/컬럼": "sample", "의미": "같은 케이스를 다른 시점에 재실행한 결과를 구분하는 태그(예: repro_20260803=이번 재실행, corebrief_prior_20260802_03=이전 실행 재채점).", "비고": "동일 채점기(고정 프롬프트)로 재채점해 v2/v3 실행별 변동폭을 직접 비교하기 위함."},
        {"영역": "기타", "지표/컬럼": "control_if/loop/try_gold", "의미": "Gold 워크플로우에 있는 if/loop/try 구조 개수(참고용 구조 통계, 채점에 직접 안 쓰임).", "비고": "-"},
        {"영역": "기타", "지표/컬럼": "unmatched_gold_count / unmatched_prediction_count", "의미": "매칭 안 된 핵심업무 액션 개수(gold=누락 FN, prediction=불필요 FP).", "비고": "Unmatched/JudgeLog 시트에서 실제 라벨 확인 가능."},
        {"영역": "생성/채점 모델", "지표/컬럼": "-", "의미": "예측 워크플로우 생성(v2/v3 에이전트)은 gpt-5.4-mini(temperature=0, seed=0). 채점 LLM 판정(핵심업무 분류/Judge Match)은 gpt-4o-mini.", "비고": "서로 다른 모델 - 생성은 더 큰 모델, 채점은 저렴한 모델로 역할 분리."},
    ]


def autosize_columns(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    for idx, column in enumerate(df.columns):
        cell_lengths = [len(str(value)) for value in df[column].tolist()]
        max_len = max([len(str(column)), *cell_lengths]) + 2
        worksheet.column_dimensions[get_column_letter(idx + 1)].width = min(max_len, 80)


def main() -> None:
    parser = argparse.ArgumentParser(description="audit_final_goldset.py의 audit.json을 상세 엑셀로 변환")
    parser.add_argument("audit_json", type=Path)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--title", default="RPA 워크플로우 생성 평가 — 핵심 지표")
    parser.add_argument("--subtitle", default=None)
    args = parser.parse_args()

    payload = load_audit(args.audit_json)
    rows = payload["results"]
    output_path = args.output or args.audit_json.with_suffix("").with_name(args.audit_json.stem + "_detail.xlsx")

    case_count = len({r["case_id"] for r in rows})
    subtitle = args.subtitle or f"확정 Goldset {case_count}개 케이스 · {payload.get('created_at', '')[:10]}"

    sheets: dict[str, pd.DataFrame] = {
        "지표설명": pd.DataFrame(build_glossary_rows()),
        "Overview": pd.DataFrame(build_overview_rows(rows)),
        "Aggregates": pd.DataFrame(build_aggregates_rows(payload["aggregates"])),
        "Matches": pd.DataFrame(build_matches_rows(rows)),
        "Unmatched": pd.DataFrame(build_unmatched_rows(rows)),
        "JudgeLog": pd.DataFrame(build_judge_log_rows(rows)),
        "CoreBusinessClassification": pd.DataFrame(build_core_business_rows(rows)),
        "BranchCoverageDetail": pd.DataFrame(build_branch_detail_rows(rows)),
    }

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            if not df.empty:
                autosize_columns(writer, name, df)
        write_headline_sheet(writer.book, payload["aggregates"], title=args.title, subtitle=subtitle)

    print(f"wrote {output_path} ({sum(len(df) for df in sheets.values())} total rows across {len(sheets) + 1} sheets)")


if __name__ == "__main__":
    main()
